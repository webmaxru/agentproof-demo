from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
WIDTH, HEIGHT, FPS = 1920, 1080, 30
BACKGROUND = "#101719"
PANEL = "#1c292c"
WHITE = "#f5f7f5"
MUTED = "#b8c8c6"
MINT = "#50ddae"
AMBER = "#ffd166"
CORAL = "#ff7f85"
PURPLE = "#bc9cff"

SEGMENTS = [
    {
        "start": 0, "end": 14, "title": "The exact change matters",
        "lowerThird": "AI-ASSISTED CHANGE - CURRENT HEAD SHA",
        "picture": "Unsafe PR, full current SHA, origin declaration, red AgentProof / gate.",
        "voiceover": "AI can help build software quickly. The release question is different: did this exact change earn evidence and an accountable decision?",
        "humanAction": "Open the prepared unsafe PR. Do not show tokens or unrelated browser chrome.",
    },
    {
        "start": 14, "end": 31, "title": "Deterministic evidence",
        "lowerThird": "READ-ONLY SPECIALISTS - SAME SHA",
        "picture": "Real check and artifact, then permitted same-SHA specialist results.",
        "voiceover": "AgentProof runs deterministic collectors without secrets and binds the result to the pull-request head SHA. Read-only specialist sessions add context for tests, security, and policy, but they remain advisory.",
        "humanAction": "Omit unavailable reviewer footage; never substitute invented reviewer output.",
    },
    {
        "start": 31, "end": 51, "title": "A board, not an approval",
        "lowerThird": "MUTABLE COORDINATION - GITHUB IS AUTHORITATIVE",
        "picture": "Evidence Board: dependency fail, required-test fail, retention unknown.",
        "voiceover": "The assembler rejects mixed repositories, policies, or SHAs. The board helps people coordinate, but it is mutable. The GitHub check, artifact, comments, reviews, and commit remain the authoritative record.",
        "humanAction": "Load only validated evidence. Keep the authority banner visible.",
    },
    {
        "start": 51, "end": 69, "title": "An explicit human decision",
        "lowerThird": "HUMAN DISPOSITION - ELIGIBLE, EXPLICIT, EXPIRING",
        "picture": "Eligible retention finding and an unsubmitted current-SHA exception command.",
        "voiceover": "A release manager may accept only a policy-eligible finding, with a specific reason and an expiry. This is not an approval, and it does not excuse the non-exceptionable failures.",
        "humanAction": "Only the authorized human may verify and submit the decision.",
    },
    {
        "start": 69, "end": 93, "title": "Remediation creates a new SHA",
        "lowerThird": "REMEDIATION COMMIT - NEW HEAD SHA",
        "picture": "Real dependency upgrade and restored authorization-test marker, then new commit.",
        "voiceover": "The preferred path is remediation. The dependency is upgraded and the authorization test is restored. The push creates a new head SHA, so this is now a new release decision.",
        "humanAction": "Advance the recording PR only after capturing the unsafe-head decision.",
    },
    {
        "start": 93, "end": 111, "title": "Old decisions do not travel",
        "lowerThird": "NEW SHA - PRIOR EVIDENCE AND APPROVAL STALE",
        "picture": "Different full head SHAs; stale disposition and fresh workflow activity.",
        "voiceover": "The old exception, evidence, and approval cannot travel across a commit change. AgentProof fails closed until the new revision earns fresh evidence.",
        "humanAction": "Show stale approval only if a real prior human approval existed and was dismissed.",
    },
    {
        "start": 111, "end": 132, "title": "Fresh evidence, same new SHA",
        "lowerThird": "FRESH SAME-SHA EVIDENCE",
        "picture": "New-PR-head check, artifact, policy and digest; fresh notes only if available.",
        "voiceover": "Fresh collectors and fresh specialist notes now describe the remediated commit. Any remaining exception is recorded again against this SHA, never copied from the previous revision.",
        "humanAction": "A human submits any still-needed exception against the new full head SHA.",
    },
    {
        "start": 132, "end": 152, "title": "Evidence plus independent review",
        "lowerThird": "GREEN GATE + INDEPENDENT REVIEW",
        "picture": "Actual green gate and artifact, then a distinct human's approval.",
        "voiceover": "The gate turns green only when the current evidence, policy, disposition, and artifact digest agree. A different human still reviews the change. Evidence alone cannot approve, and approval alone cannot bypass a red gate.",
        "humanAction": "Requires enforced repository protections. Show merge available, but never merge.",
    },
    {
        "start": 152, "end": 163, "title": "Bounded evidence",
        "lowerThird": "BOUNDED TO REPOSITORY + POLICY + SHA",
        "picture": "Limits frame: bounded evidence, no universal provenance, no automatic approval.",
        "voiceover": "This is bounded evidence for one repository, policy, and commit. It does not prove universal authorship, legal compliance, or production suitability.",
        "humanAction": "Use limits slide 3. Do not imply production or legal certification.",
    },
    {
        "start": 163, "end": 174, "title": "Stop at the permission boundary",
        "lowerThird": "PERMISSION CANARY - DO NOT RE-ENABLE",
        "picture": "Actual disabled canary with UNSAFE_TOOL_BOUNDARY.\nOtherwise show a labeled limitation card.",
        "voiceover": "We tested the automation boundary too. Mutation tools remained visible, so AgentProof stopped before using them and the trigger was disabled. Refusing an unsafe runtime is part of the product.",
        "humanAction": "Use this wording only with an actual retained result and verified disabled state.",
    },
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    filename = "segoeuib.ttf" if bold else "segoeui.ttf"
    candidates = [
        Path("C:\\Windows\\Fonts") / filename,
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold
             else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    raise FileNotFoundError("Install Segoe UI or DejaVu Sans before rendering media.")


def wrapped(draw: ImageDraw.ImageDraw, text: str, x: int, y: int, width: int,
            size: int, color: str = WHITE, bold: bool = False, gap: int = 14) -> int:
    face = font(size, bold)
    lines: list[str] = []
    for paragraph in text.split("\n"):
        current = ""
        for word in paragraph.split():
            candidate = f"{current} {word}".strip()
            if draw.textlength(candidate, font=face) > width and current:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    for line in lines:
        draw.text((x, y), line, font=face, fill=color)
        y += size + gap
    return y


def stamp(seconds: float) -> str:
    milliseconds = round(seconds * 1000)
    hours, milliseconds = divmod(milliseconds, 3_600_000)
    minutes, milliseconds = divmod(milliseconds, 60_000)
    secs, milliseconds = divmod(milliseconds, 1000)
    return f"{hours:02}:{minutes:02}:{secs:02},{milliseconds:03}"


def build_captions() -> None:
    def split_sentence(sentence: str) -> list[str]:
        lines = textwrap.wrap(sentence, width=46, break_long_words=False,
                              break_on_hyphens=False)
        words = sentence.split()
        if len(lines) <= 2 and len(words) <= 14:
            return [sentence]
        candidates = list(range(4, len(words) - 3)) or [len(words) // 2]
        split_at = min(candidates, key=lambda index:
                       abs(index - len(words) / 2) -
                       (3 if words[index - 1].endswith((",", ":", ";")) else 0))
        return split_sentence(" ".join(words[:split_at])) + split_sentence(" ".join(words[split_at:]))

    cues = []
    for segment in SEGMENTS:
        sentences = re.split(r"(?<=[.!?])\s+", segment["voiceover"])
        chunks = [chunk for sentence in sentences for chunk in split_sentence(sentence)]
        word_count = sum(len(chunk.split()) for chunk in chunks)
        cursor = segment["start"] + 0.5
        available = min(segment["end"] - segment["start"] - 1.0, word_count / (140 / 60))
        for chunk in chunks:
            duration = available * len(chunk.split()) / word_count
            formatted = textwrap.fill(chunk, width=46, break_long_words=False,
                                      break_on_hyphens=False)
            assert len(formatted.splitlines()) <= 2
            assert 1 <= duration <= 7
            cues.append((cursor, cursor + duration, formatted))
            cursor += duration
    assert cues[-1][1] <= 174
    body = "\n\n".join(
        f"{index}\n{stamp(start)} --> {stamp(end)}\n{text}"
        for index, (start, end, text) in enumerate(cues, 1)
    )
    (ROOT / "agentproof-voiceover-draft.srt").write_text(
        body + "\n", encoding="utf-8", newline="\n"
    )


def build_images() -> None:
    stills = ROOT / "cue-cards"
    overlays = ROOT / "overlays"
    stills.mkdir(exist_ok=True)
    overlays.mkdir(exist_ok=True)
    for index, segment in enumerate(SEGMENTS):
        image = Image.new("RGB", (WIDTH, HEIGHT), BACKGROUND)
        draw = ImageDraw.Draw(image)
        draw.rectangle((0, 0, WIDTH, 90), fill=AMBER)
        draw.text((90, 20), "PRECOMPUTED / NOT LIVE", font=font(42, True), fill=BACKGROUND)
        draw.text((90, 165), f"AGENTPROOF   |   SHOT {index + 1:02} / 10",
                  font=font(32, True), fill=MINT)
        wrapped(draw, segment["title"], 90, 255, 1740, 78, bold=True)
        draw.rounded_rectangle((90, 455, 1830, 805), radius=24, fill=PANEL)
        draw.text((130, 487), "REQUIRED LIVE PICTURE", font=font(28, True), fill=MINT)
        wrapped(draw, segment["picture"], 130, 550, 1650, 48)
        wrapped(draw, segment["humanAction"], 90, 855, 1740, 31, MUTED)
        label = f"{segment['start'] // 60}:{segment['start'] % 60:02} - {segment['end'] // 60}:{segment['end'] % 60:02}"
        draw.text((90, 1000), label, font=font(28, True), fill=WHITE)
        draw.text((550, 1000), "TIMING REFERENCE ONLY. REPLACE WITH REAL FOOTAGE.",
                  font=font(28), fill=MUTED)
        image.save(stills / f"{index:02}.png", optimize=True)

        overlay = Image.new("RGBA", (WIDTH, HEIGHT), (0, 0, 0, 0))
        overlay_draw = ImageDraw.Draw(overlay)
        overlay_draw.rounded_rectangle((80, 900, 1840, 1020), radius=14, fill=(16, 23, 25, 242))
        overlay_draw.rectangle((80, 900, 94, 1020), fill=MINT)
        overlay_draw.text((124, 929), segment["lowerThird"], font=font(35, True), fill=WHITE)
        overlay.save(overlays / f"{index:02}-lower-third.png", optimize=True)

    banner = Image.new("RGBA", (WIDTH, HEIGHT), (0, 0, 0, 0))
    draw = ImageDraw.Draw(banner)
    draw.rectangle((0, 0, WIDTH, 90), fill=AMBER)
    draw.text((90, 20), "PRECOMPUTED / NOT LIVE", font=font(42, True), fill=BACKGROUND)
    banner.save(overlays / "precomputed-not-live.png", optimize=True)


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#"))


def text(slide, value: str, x: float, y: float, width: float, height: float,
         size: int, color: str = WHITE, bold: bool = False, name: str = "") -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    shape.name = name or value[:45]
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0)
    frame.margin_top = frame.margin_bottom = Inches(0)
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(6)
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Segoe UI"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)


def box(slide, x: float, y: float, width: float, height: float, color: str = PANEL) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()


def base_slide(presentation: Presentation, index: int, eyebrow: str):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BACKGROUND)
    text(slide, f"AGENTPROOF  /  {eyebrow}", .65, .45, 12, .4, 16, MINT, True)
    text(slide, "ILLUSTRATIVE / NOT LIVE EVIDENCE", .65, 6.95, 10.8, .35, 16, MUTED)
    text(slide, f"0{index}", 12, 6.95, .65, .35, 16, MINT, True)
    return slide


def build_deck() -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "AgentProof - Recording support"
    presentation.core_properties.subject = "Illustrative slides, not live scanner evidence"
    presentation.core_properties.author = "AgentProof contributors"

    slide = base_slide(presentation, 1, "THE TRUST GAP")
    text(slide, "Agents can build.\nPeople decide what ships.", .65, 1.05, 12, 1.6, 42, bold=True)
    text(slide, "The question is whether this exact change earned evidence.", .65, 2.85,
         12, .65, 24, MUTED)
    for x, label, heading, body, accent in [
        (.65, "FAIL", "Dependency evidence", "Runtime dependency exceeds\nthe protected severity limit.", CORAL),
        (4.75, "FAIL", "Required behavior", "Required test ID is missing.\nA passing suite is not enough.", CORAL),
        (8.85, "UNKNOWN", "Retention declaration", "A missing deletion fact is\nnot a positive result.", AMBER),
    ]:
        box(slide, x, 3.95, 3.8, 2.55)
        text(slide, label, x + .22, 4.2, 3.35, .35, 18, accent, True)
        text(slide, heading, x + .22, 4.75, 3.35, .55, 23, bold=True)
        text(slide, body, x + .22, 5.4, 3.35, .85, 16, MUTED)
    slide.notes_slide.notes_text_frame.text = (
        "Use this as an illustrative supporting slide, not the first frame of the live video. "
        "The opening must show the actual unsafe PR and check. These are expected synthetic "
        "findings, not fresh scanner output. No universal provenance or certification claim."
    )

    slide = base_slide(presentation, 2, "THE CONTROL LOOP")
    text(slide, "Evidence follows the commit.", .65, 1.05, 12, .85, 42, bold=True)
    for x, title, detail in [
        (.65, "1  COLLECT", "Exact PR head SHA\nNo secrets; read-only"),
        (4.9, "2  EVALUATE", "Protected-base policy\nDeterministic gate"),
        (9.15, "3  PUBLISH", "GitHub check + artifact\nCurrent-SHA summary"),
    ]:
        box(slide, x, 2.25, 3.5, 1.65)
        text(slide, title, x + .22, 2.48, 3.08, .4, 21, MINT, True)
        text(slide, detail, x + .22, 3.08, 3.08, .7, 17, WHITE)
    text(slide, ">", 4.34, 2.88, .38, .5, 28, MINT, True)
    text(slide, ">", 8.59, 2.88, .38, .5, 28, MINT, True)
    box(slide, .65, 4.4, 7.45, 1.45)
    text(slide, "MANUAL, READ-ONLY SPECIALISTS", .9, 4.6, 6.95, .35, 19, PURPLE, True)
    text(slide, "Test  |  Security  |  Policy  ->  Mutable Evidence Board",
         .9, 5.2, 6.95, .35, 17, WHITE)
    text(slide, ">", 8.18, 4.93, .3, .5, 24, MINT, True)
    box(slide, 8.5, 4.4, 4.15, 1.45)
    text(slide, "HUMAN DECISIONS", 8.75, 4.6, 3.65, .35, 19, AMBER, True)
    text(slide, "Eligible exception\nIndependent code review",
         8.75, 5.13, 3.65, .65, 17, WHITE)
    text(slide, "NEW COMMIT  ->  OLD EVIDENCE AND DECISIONS ARE STALE",
         .65, 6.25, 12, .45, 22, MINT, True)
    slide.notes_slide.notes_text_frame.text = (
        "GitHub records and enforced repository protections are the authority. "
        "Reviewer prose cannot make the gate more favorable. The three specialists and "
        "assembler are manual. A new head invalidates prior evidence and dispositions. "
        "Only demonstrate protected merge behavior after repository entitlement and "
        "the actual ruleset are verified."
    )

    slide = base_slide(presentation, 3, "HONEST BOUNDARIES")
    text(slide, "Bounded evidence.\nAccountable decisions.", .65, 1.05, 8.5, 1.65, 42, bold=True)
    for y, number, title, detail in [
        (3.15, "01", "No universal provenance", "Origin can be self-declared or unknown."),
        (4.18, "02", "No automatic approval", "Humans accept eligible exceptions and review code."),
        (5.21, "03", "No immutable canvas", "GitHub checks, comments, reviews and artifacts are authoritative."),
    ]:
        text(slide, number, .65, y, .6, .45, 24, MINT, True)
        text(slide, title, 1.45, y, 6.5, .45, 24, bold=True)
        text(slide, detail, 1.45, y + .48, 6.6, .5, 16, MUTED)
    box(slide, 8.7, 2.98, 3.95, 3.45)
    text(slide, "PERMISSION BOUNDARY", 8.95, 3.25, 3.45, .35, 16, AMBER, True)
    text(slide, "Mutation tools?\nSTOP.", 8.95, 3.93, 3.45, 1.05, 32, CORAL, True)
    text(slide, "Use only approved read tools.\nKeep a failed canary disabled.",
         8.95, 5.2, 3.45, .85, 16, MUTED)
    slide.notes_slide.notes_text_frame.text = (
        "This slide is safe for the limits segment. It is an explanation, not evidence "
        "that a canary ran in this repository. Show UNSAFE_TOOL_BOUNDARY and a disabled "
        "state only from an actual verified run. If unavailable, use the explicitly "
        "labeled fallback and say that runtime validation remains blocked. "
        "Do not call the result legally compliant, universally secure, or approved."
    )
    presentation.save(ROOT / "agentproof-recording-deck.pptx")


def build_video() -> None:
    commands = []
    for index, segment in enumerate(SEGMENTS):
        commands.append(f"file 'cue-cards/{index:02}.png'")
        commands.append(f"duration {segment['end'] - segment['start']}")
    commands.append("file 'cue-cards/09.png'")
    (ROOT / "cue-cards.ffconcat").write_text(
        "ffconcat version 1.0\n" + "\n".join(commands) + "\n",
        encoding="utf-8", newline="\n"
    )
    subprocess.run([
        "ffmpeg", "-hide_banner", "-loglevel", "error", "-y",
        "-f", "concat", "-safe", "0", "-i", str(ROOT / "cue-cards.ffconcat"),
        "-t", "174", "-r", str(FPS), "-c:v", "libx264", "-preset", "medium",
        "-tune", "stillimage", "-crf", "22", "-pix_fmt", "yuv420p",
        "-movflags", "+faststart", "-an", str(ROOT / "agentproof-cue-timing-draft.mp4"),
    ], check=True, cwd=ROOT)


def verify() -> dict:
    assert len(SEGMENTS) == 10
    assert SEGMENTS[0]["start"] == 0 and SEGMENTS[-1]["end"] == 174
    for current, following in zip(SEGMENTS, SEGMENTS[1:]):
        assert current["end"] == following["start"]
    presentation = Presentation(ROOT / "agentproof-recording-deck.pptx")
    assert len(presentation.slides) == 3
    assert abs(presentation.slide_width / presentation.slide_height - 16 / 9) < 0.001
    for index, slide in enumerate(presentation.slides):
        assert "ILLUSTRATIVE / NOT LIVE EVIDENCE" in "\n".join(
            shape.text for shape in slide.shapes if shape.has_text_frame
        )
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, (index, shape.name)
            assert shape.left + shape.width <= presentation.slide_width, (index, shape.name)
            assert shape.top + shape.height <= presentation.slide_height, (index, shape.name)
    for path in (ROOT / "cue-cards").glob("*.png"):
        with Image.open(path) as image:
            assert image.size == (WIDTH, HEIGHT)
            assert image.getpixel((10, 10)) == (255, 209, 102)
    for path in (ROOT / "overlays").glob("*.png"):
        with Image.open(path) as image:
            assert image.size == (WIDTH, HEIGHT) and image.mode == "RGBA"
    forbidden = "msft-common" + "-demos"
    for path in ROOT.glob("*"):
        if path.suffix in {".md", ".json", ".srt", ".txt"}:
            assert forbidden not in path.read_text(encoding="utf-8"), path
    video = ROOT / "agentproof-cue-timing-draft.mp4"
    video_metadata = None
    if video.exists():
        probe = subprocess.run([
            "ffprobe", "-v", "error", "-show_entries",
            "stream=codec_type,width,height,r_frame_rate:format=duration",
            "-of", "json", str(video),
        ], capture_output=True, text=True, check=True)
        video_metadata = json.loads(probe.stdout)
        stream = video_metadata["streams"][0]
        assert stream["width"] == WIDTH and stream["height"] == HEIGHT
        assert stream["r_frame_rate"] == "30/1"
        assert abs(float(video_metadata["format"]["duration"]) - 174) <= 1 / FPS
    files = []
    for path in sorted(ROOT.rglob("*")):
        if path.is_file() and path.name != "asset-manifest.json":
            data = path.read_bytes()
            if path.suffix in {".md", ".json", ".mjs", ".py", ".srt", ".ffconcat", ".txt"}:
                assert b"\r\n" not in data, f"Use LF before hashing text asset: {path.name}"
            files.append({
                "path": path.relative_to(ROOT).as_posix(),
                "bytes": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
            })
    manifest = {
        "status": "PRECOMPUTED / NOT LIVE",
        "purpose": "Recording preparation; no live footage, approval or privacy sign-off",
        "width": WIDTH, "height": HEIGHT, "fps": FPS, "durationSeconds": 174,
        "slideCount": 3, "voiceoverWordCount": sum(len(s["voiceover"].split()) for s in SEGMENTS),
        "video": video_metadata, "files": files,
    }
    (ROOT / "asset-manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--video", action="store_true")
    parser.add_argument("--verify-only", action="store_true")
    args = parser.parse_args()
    if not args.verify_only:
        (ROOT / "timeline.json").write_text(
            json.dumps(SEGMENTS, indent=2) + "\n", encoding="utf-8", newline="\n"
        )
        build_captions()
        build_images()
        build_deck()
        if args.video:
            build_video()
    manifest = verify()
    print(json.dumps({key: manifest[key] for key in
                      ("status", "durationSeconds", "slideCount", "voiceoverWordCount")}))


if __name__ == "__main__":
    main()
